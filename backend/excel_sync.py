import os
import shutil
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger("workwise.excel_sync")

EXCEL_COLUMNS = [
    "event_id",
    "event_type",
    "task_id",
    "task_title",
    "employee_id",
    "employee_name",
    "previous_employee_id",
    "previous_employee_name",
    "priority",
    "urgency",
    "deadline",
    "sla",
    "location",
    "estimated_hours",
    "remaining_hours_before",
    "remaining_hours_after",
    "matching_score_or_rank",
    "constraint_checks",
    "trigger_reason",
    "actor",
    "timestamp",
    "model_version",
    "notes"
]

COLUMN_DESCRIPTIONS = {
    "event_id": ("UUID or unique identifier for the history event", "Text"),
    "event_type": ("Type of event: assignment_created, approved, rejected, reassigned, completed, cancelled, reevaluated", "Text"),
    "task_id": ("Unique task ID reference", "Integer"),
    "task_title": ("Human-readable title of the task at event time", "Text"),
    "employee_id": ("Unique employee ID reference", "Integer"),
    "employee_name": ("Full name of assigned employee at event time", "Text"),
    "previous_employee_id": ("Previous employee ID (for reassignment events)", "Integer/Empty"),
    "previous_employee_name": ("Previous employee name (for reassignment events)", "Text/Empty"),
    "priority": ("Task priority level (Low, Medium, High, Critical)", "Text"),
    "urgency": ("Task urgency level (Low, Medium, High, Urgent)", "Text"),
    "deadline": ("Task deadline timestamp (ISO 8601)", "Text"),
    "sla": ("Service Level Agreement commitment in hours", "Numeric (Hours)"),
    "location": ("Geographic location or department relevant to assignment", "Text"),
    "estimated_hours": ("Total estimated effort required for task", "Numeric (Hours)"),
    "remaining_hours_before": ("Employee's remaining workload capacity before this event", "Numeric (Hours)"),
    "remaining_hours_after": ("Employee's remaining workload capacity after this event", "Numeric (Hours)"),
    "matching_score_or_rank": ("Random Forest model suitability rank/score output", "Text/Numeric"),
    "constraint_checks": ("Summary of hard constraint evaluation (passed/failed rules)", "Text"),
    "trigger_reason": ("Trigger rationale for the action (manager approval, availability change, etc.)", "Text"),
    "actor": ("User, manager, or system component initiating the event", "Text"),
    "timestamp": ("Event occurrence time (ISO 8601)", "Text"),
    "model_version": ("Random Forest model version used for decision", "Text"),
    "notes": ("Detailed rationale, rejection reason, or system diagnostics", "Text")
}

def get_default_excel_path() -> Path:
    base_dir = Path(__file__).resolve().parent.parent / "data" / "history"
    base_dir.mkdir(parents=True, exist_ok=True)
    return base_dir / "allocation_history.xlsx"

def create_readme_sheet(wb: openpyxl.Workbook):
    ws = wb.active
    ws.title = "README"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sub_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws.cell(row=1, column=1, value="WorkWise AI — Shared Allocation History Log Documentation").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.merge_cells("A1:C1")
    ws.row_dimensions[1].height = 35

    ws.cell(row=3, column=1, value="Overview:").font = bold_font
    ws.cell(row=3, column=2, value="This shared workbook logs every workload allocation and decision event produced by WorkWise AI. SQLite remains the transactional source of truth, while this workbook provides team members with a readable audit mirror.").font = Font(name="Calibri", size=11)

    ws.cell(row=5, column=1, value="Column Name").font = header_font
    ws.cell(row=5, column=1).fill = sub_fill
    ws.cell(row=5, column=2, value="Description & Purpose").font = header_font
    ws.cell(row=5, column=2).fill = sub_fill
    ws.cell(row=5, column=3, value="Data Type / Format").font = header_font
    ws.cell(row=5, column=3).fill = sub_fill
    ws.row_dimensions[5].height = 25

    row_idx = 6
    for col_name in EXCEL_COLUMNS:
        desc, dtype = COLUMN_DESCRIPTIONS.get(col_name, ("N/A", "Text"))
        c1 = ws.cell(row=row_idx, column=1, value=col_name)
        c2 = ws.cell(row=row_idx, column=2, value=desc)
        c3 = ws.cell(row=row_idx, column=3, value=dtype)
        c1.font = bold_font
        for c in [c1, c2, c3]:
            c.border = border
        row_idx += 1

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 85
    ws.column_dimensions['C'].width = 25

def format_history_sheet(ws: openpyxl.worksheet.worksheet.Worksheet):
    ws.views.sheetView[0].showGridLines = True
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

def _ensure_workbook(excel_path: Path) -> openpyxl.Workbook:
    if excel_path.exists():
        wb = openpyxl.load_workbook(excel_path)
    else:
        wb = openpyxl.Workbook()
        create_readme_sheet(wb)
        ws_hist = wb.create_sheet(title="Allocation History")
        format_history_sheet(ws_hist)
    return wb

def append_history_row(event: Dict[str, Any], excel_path_str: Optional[str] = None, db_conn=None) -> bool:
    """
    Appends a single history row to allocation_history.xlsx safely.
    Excel sync failures will be caught and logged without breaking SQLite transactions.
    """
    path = Path(excel_path_str) if excel_path_str else get_default_excel_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.parent / f"temp_{path.name}"

    try:
        wb = _ensure_workbook(path)
        if "Allocation History" in wb.sheetnames:
            ws = wb["Allocation History"]
        else:
            ws = wb.create_sheet(title="Allocation History")
            format_history_sheet(ws)

        format_history_sheet(ws)

        row_values = []
        for col in EXCEL_COLUMNS:
            val = event.get(col, "")
            if val is None:
                val = ""
            row_values.append(str(val) if not isinstance(val, (int, float)) else val)

        ws.append(row_values)

        for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
            col_letter = get_column_letter(col_idx)
            current_w = ws.column_dimensions[col_letter].width or 12
            ws.column_dimensions[col_letter].width = max(current_w, len(col_name) + 4)

        wb.save(temp_path)
        wb.close()
        shutil.move(str(temp_path), str(path))

        _update_settings_sync_status(path, True, "OK", conn=db_conn)
        return True

    except Exception as e:
        logger.warning(f"Failed to append row to Excel history log at {path}: {str(e)}")
        if temp_path.exists():
            try:
                temp_path.unlink()
            except Exception:
                pass
        _update_settings_sync_status(path, False, f"Sync Failed: {str(e)}", conn=db_conn)
        return False

def rebuild_excel_from_sqlite(conn, excel_path_str: Optional[str] = None) -> Dict[str, Any]:
    path = Path(excel_path_str) if excel_path_str else get_default_excel_path()
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        backup_path = path.parent / f"backup_{path.name}"
        try:
            shutil.copy2(str(path), str(backup_path))
        except Exception as e:
            logger.warning(f"Failed to create Excel backup: {str(e)}")

    cursor = conn.cursor()
    cursor.execute("SELECT * FROM history ORDER BY id ASC")
    rows = cursor.fetchall()

    wb = openpyxl.Workbook()
    create_readme_sheet(wb)
    ws_hist = wb.create_sheet(title="Allocation History")
    format_history_sheet(ws_hist)

    border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    for r_idx, row in enumerate(rows, 2):
        row_dict = dict(row)
        row_values = []
        for col in EXCEL_COLUMNS:
            val = row_dict.get(col, "")
            if val is None:
                val = ""
            row_values.append(val)
        ws_hist.append(row_values)
        for c_idx in range(1, len(EXCEL_COLUMNS) + 1):
            ws_hist.cell(row=r_idx, column=c_idx).border = border

    for col_idx, col_name in enumerate(EXCEL_COLUMNS, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(col_name)
        for row in range(1, len(rows) + 2):
            cell_val = ws_hist.cell(row=row, column=col_idx).value
            if cell_val:
                max_len = max(max_len, min(len(str(cell_val)), 50))
        ws_hist.column_dimensions[col_letter].width = max(max_len + 3, 12)

    temp_path = path.parent / f"temp_{path.name}"
    wb.save(temp_path)
    wb.close()
    shutil.move(str(temp_path), str(path))

    now_str = datetime.now().isoformat()
    cursor.execute("""
        INSERT OR REPLACE INTO settings (key, value, updated_at)
        VALUES ('excel_log_path', ?, datetime('now')),
               ('last_excel_sync', ?, datetime('now')),
               ('excel_sync_status', 'OK', datetime('now'))
    """, (str(path), now_str))
    conn.commit()

    return {
        "status": "success",
        "rows_written": len(rows),
        "excel_path": str(path),
        "timestamp": now_str
    }

def _update_settings_sync_status(path: Path, success: bool, message: str, conn=None):
    try:
        close_conn = False
        if conn is None:
            from backend.db import get_db
            conn = get_db()
            close_conn = True
        cursor = conn.cursor()
        now_str = datetime.now().isoformat()
        status_val = "OK" if success else message
        cursor.execute("INSERT OR REPLACE INTO settings (key, value, updated_at) VALUES ('excel_log_path', ?, datetime('now'))", (str(path),))
        if success:
            cursor.execute("INSERT OR REPLACE INTO settings (key, value, updated_at) VALUES ('last_excel_sync', ?, datetime('now'))", (now_str,))
        cursor.execute("INSERT OR REPLACE INTO settings (key, value, updated_at) VALUES ('excel_sync_status', ?, datetime('now'))", (status_val,))
        conn.commit()
        if close_conn:
            conn.close()
    except Exception as e:
        logger.debug(f"Could not update settings sync status: {str(e)}")
