import os
import pytest
import openpyxl
from pathlib import Path
from backend.db import init_db, get_db
from backend.excel_sync import append_history_row, rebuild_excel_from_sqlite, EXCEL_COLUMNS

@pytest.fixture
def test_env(tmp_path, monkeypatch):
    db_path = tmp_path / "test_excel.db"
    excel_path = tmp_path / "allocation_history_test.xlsx"

    monkeypatch.setattr("backend.db.DB_PATH", db_path)
    monkeypatch.setattr("backend.excel_sync.get_default_excel_path", lambda: excel_path)

    init_db()
    return db_path, excel_path

def test_excel_append_and_rebuild(test_env):
    db_path, excel_path = test_env

    # 1. Test append_history_row
    event = {
        "event_id": "EVT-TEST001",
        "event_type": "assignment_created",
        "task_id": 1,
        "task_title": "Build Test Auth API",
        "employee_id": 10,
        "employee_name": "Alice Developer",
        "previous_employee_id": None,
        "previous_employee_name": None,
        "priority": "High",
        "urgency": "High",
        "deadline": "2026-10-01T17:00:00",
        "sla": "24",
        "location": "New York",
        "estimated_hours": 10.0,
        "remaining_hours_before": 40.0,
        "remaining_hours_after": 30.0,
        "matching_score_or_rank": "Rank #1 (Score: 0.9500)",
        "constraint_checks": "Passed",
        "trigger_reason": "Manager approval",
        "actor": "Manager",
        "timestamp": "2026-09-18T10:00:00",
        "model_version": "v1.0.0",
        "notes": "Initial assignment"
    }

    success = append_history_row(event, str(excel_path))
    assert success is True
    assert excel_path.exists()

    # Load workbook and check sheet structure
    wb = openpyxl.load_workbook(excel_path)
    assert "README" in wb.sheetnames
    assert "Allocation History" in wb.sheetnames

    ws_hist = wb["Allocation History"]
    # Header row + 1 data row
    assert ws_hist.max_row == 2

    headers = [ws_hist.cell(row=1, column=c).value for c in range(1, len(EXCEL_COLUMNS) + 1)]
    assert headers == EXCEL_COLUMNS

    # Verify first data row
    row1_event_id = ws_hist.cell(row=2, column=1).value
    assert row1_event_id == "EVT-TEST001"

    # 2. Test rebuild_excel_from_sqlite
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO history (
            event_id, event_type, task_id, task_title, employee_id, employee_name,
            priority, urgency, deadline, sla, location, estimated_hours,
            remaining_hours_before, remaining_hours_after, matching_score_or_rank,
            constraint_checks, trigger_reason, actor, timestamp, model_version, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        "EVT-TEST002", "completed", 1, "Build Test Auth API", 10, "Alice Developer",
        "High", "High", "2026-10-01T17:00:00", "24", "New York", 10.0,
        30.0, 40.0, "N/A", "Passed", "User completed task", "Manager",
        "2026-09-18T12:00:00", "v1.0.0", "Task completed"
    ))
    conn.commit()

    rebuild_res = rebuild_excel_from_sqlite(conn, str(excel_path))
    conn.close()

    assert rebuild_res["status"] == "success"
    assert rebuild_res["rows_written"] == 1

    wb_rebuilt = openpyxl.load_workbook(excel_path)
    ws_rebuilt = wb_rebuilt["Allocation History"]
    assert ws_rebuilt.cell(row=2, column=1).value == "EVT-TEST002"
